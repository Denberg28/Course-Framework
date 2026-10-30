import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from course_core import (CourseError, SAMPLE, MAX_FILE_BYTES, analyze_syllabus, atomic_write,
                         build_course_markdown, export_bundle, export_html, extract_text,
                         course_allocation, load_project, make_project, project_bytes)


class CourseTests(unittest.TestCase):
    def test_ae417_table_keeps_thirteen_modules_and_supplied_bibliography(self):
        source = Path('examples/AE417-syllabus.md').read_text(encoding='utf-8')
        plan = analyze_syllabus(source)
        self.assertEqual(plan['metadata']['code'], 'AE 417')
        self.assertEqual(plan['metadata']['title'], 'AERONAUTICAL ENGINEERING PRACTIVE (OJT) LECTURE')
        self.assertTrue(plan['metadata']['description'].startswith('The course consists of a preparatory'))
        self.assertEqual(plan['metadata']['units'], '3 units lecture, 0 unit laboratory = 3 Units')
        self.assertEqual(plan['metadata']['weekly_hours'], '3 hours lecture, 0 hours laboratory = 3 Hours')
        self.assertEqual(plan['metadata']['prerequisite'], 'Fourth Year Standing')
        self.assertEqual((plan['metadata']['lecture_hours'], plan['metadata']['lab_hours']), (54, 0))
        self.assertEqual(len(plan['outcomes']), 7)
        self.assertEqual([t['title'] for t in plan['topics']], [
            'Internship Orientation and Policies', 'Aviation Industry Overview / Careers / Making Resume',
            'CAAP Regulations and Airworthiness Requirements', 'Aviation Safety Management System (SMS)',
            'Quality Assurance and Quality Control Systems', 'Aircraft Maintenance Documentation and Records',
            'Engineering Ethics and Professional Responsibility', 'Human Factors in Aviation Maintenance',
            'Occupational Safety and Health Standards', 'Technical Report Writing and Logbook Preparation',
            'Research Application in Aviation Industry', 'Career Development and Professional Practice',
            'Pre-Deployment Seminar and Internship Requirements'])
        self.assertEqual([t['lecture_hours'] for t in plan['topics']], [3, 3, 6, 6, 6, 6, 3, 3, 3, 6, 3, 3, 3])
        # Empty/dash cells are unspecified per module, not inferred numeric zeroes.
        self.assertTrue(all(not t['subtopics'] and 'lab_hours' not in t for t in plan['topics']))
        self.assertEqual(len(plan['references']), 8)
        self.assertTrue(plan['references'][0].startswith('Kroes,'))
        self.assertIn('CAAP', plan['references'][4])
        self.assertFalse(plan['warnings'])

    def test_ae417_clipboard_spaces_entities_and_word_table_match_markdown(self):
        source = Path('examples/AE417-syllabus.md').read_text(encoding='utf-8')
        expected = analyze_syllabus(source)
        rows = [[c.strip() for c in line[1:-1].split('|')] for line in source.splitlines()
                if not line.startswith('|---')]
        for separator in ('\t', '  '):
            with self.subTest(separator=separator):
                self.assertEqual(analyze_syllabus('\n'.join(separator.join(r) for r in rows)), expected)
        decorated = source.replace(' | ', '\u00a0 | \u00a0').replace('<br>', '<br />')
        decorated = decorated.replace('Course Code:', '**Course   Code:**').replace(' | |', ' | &#x20; |')
        self.assertEqual(analyze_syllabus(decorated), expected)
        from docx import Document
        document = Document()
        table = document.add_table(rows=len(rows), cols=4)
        for row, values in zip(table.rows, rows):
            for cell, text in zip(row.cells, values):
                cell.text = text.replace('<br>', '\n')
        target = io.BytesIO()
        document.save(target)
        self.assertEqual(analyze_syllabus(extract_text('ojt.docx', target.getvalue())), expected)

    def test_hour_cells_stay_in_their_columns_when_empty_or_reversed(self):
        source = ('Course Name: Table\n| Topic | LAB HRS | LEC HRS |\n|---|---|---|\n'
                  '| 1. First | | 3 |\n| 2. Second | 2.5 | |\n| 3. Third | – | 6 |\n'
                  '| 4. Fourth | — | 3 |\nReferences:\nSupplied source')
        plan = analyze_syllabus(source)
        self.assertEqual([(t.get('lecture_hours'), t.get('lab_hours')) for t in plan['topics']],
                         [(3, None), (None, 2.5), (6, None), (3, None)])
        self.assertTrue(any('no lecture hours' in w for w in plan['warnings']))
        self.assertTrue(any('no lab hours' in w for w in plan['warnings']))

    def test_full_ae416_syllabus_preserves_metadata_outcomes_readings_and_hours(self):
        source = Path('examples/AE416-syllabus.txt').read_text(encoding='utf-8')
        plan = analyze_syllabus(source)
        self.assertEqual(plan['metadata']['code'], 'AE 416')
        self.assertEqual(plan['metadata']['title'], 'RESEARCH METHODS AND APPLICATION')
        self.assertTrue(plan['metadata']['description'].startswith('The course includes lectures'))
        self.assertEqual(plan['metadata']['units'], '2 units lecture, 1 unit laboratory = 3 Units')
        self.assertEqual(plan['metadata']['weekly_hours'], '2 hours lecture, 3 hours laboratory = 5 Hours')
        self.assertEqual(plan['metadata']['prerequisite'], 'Fourth Year Standing')
        self.assertEqual(plan['metadata']['hours'], '90')
        self.assertEqual(len(plan['outcomes']), 3)
        self.assertEqual([t['title'] for t in plan['topics']], [
            'Nature and Characteristics of Research', 'Research Problem and Objectives',
            'Review of Related Literature', 'Research Design',
            'Selection and Finalization of Research Topic Applicable to Aviation Industry'])
        self.assertEqual([(t['lecture_hours'], t['lab_hours']) for t in plan['topics']],
                         [(6, 6), (6, 9), (8, 12), (8, 12), (8, 15)])
        self.assertTrue(all(not t['schedule'] for t in plan['topics']))
        self.assertEqual(len(plan['references']), 3)
        self.assertTrue(plan['references'][1].startswith('Adanza,'))
        self.assertFalse(plan['warnings'])

    def test_pasted_hour_table_tabs_spaces_and_pipes_agree(self):
        source = Path('examples/AE416-syllabus.txt').read_text(encoding='utf-8')
        expected = analyze_syllabus(source)
        for separator in (' ', '    ', ' | '):
            with self.subTest(separator=separator):
                self.assertEqual(analyze_syllabus(source.replace('\t', separator)), expected)

    def test_hour_allocation_mismatch_is_flagged_without_rewriting_source(self):
        source = 'Course Name: Research\nLearning Outcomes:\n1. Evaluate evidence.\nLEC HRS LAB HRS\n1. Methods 6.5 9\nTotal: 8 9\nTextbook\nSupplied reading'
        plan = analyze_syllabus(source)
        self.assertEqual(plan['metadata']['lecture_hours'], 8)
        self.assertEqual(plan['topics'][0]['lecture_hours'], 6.5)
        self.assertEqual(len(plan['warnings']), 1)
        self.assertIn('differs', plan['warnings'][0])
        self.assertIn('6.5 | 9', course_allocation(plan))
        reverse = analyze_syllabus(source.replace('LEC HRS LAB HRS', 'LAB HRS LEC HRS'))
        self.assertEqual(reverse['topics'][0]['lab_hours'], 6.5)

    def test_metadata_outcomes_and_references_are_not_modules(self):
        plan = analyze_syllabus(SAMPLE)
        self.assertEqual(plan['metadata']['title'], 'Aircraft Structures')
        self.assertEqual(len(plan['topics']), 5)
        self.assertEqual(len(plan['outcomes']), 3)
        self.assertEqual(plan['topics'][2]['schedule'], '3-4')
        self.assertEqual(plan['topics'][0]['subtopics'], ['Wings, fuselage, and empennage'])
        self.assertEqual(len(plan['references']), 1)

    def test_table_syllabus(self):
        text = '''| Field | Details |
|---|---|
| Course Name | Flight Theory |
Course Outline:
| Week | Topics | Hours |
|---|---|---|
| 1 | Lift and drag | 3 |
| 2-3 | Stability | 6 |
| | Total | 9 |'''
        plan = analyze_syllabus(text)
        self.assertEqual(plan['metadata']['title'], 'Flight Theory')
        self.assertEqual([t['title'] for t in plan['topics']], ['Lift and drag', 'Stability'])

    def test_plain_outline_and_empty_input(self):
        plan = analyze_syllabus('Introduction\nAircraft loads\nInspection')
        self.assertEqual(len(plan['topics']), 3)
        self.assertTrue(plan['warnings'])
        with self.assertRaises(CourseError):
            analyze_syllabus('  ')
        self.assertEqual(len(analyze_syllabus('- Introduction\n- Aircraft loads\n- Inspection')['topics']), 3)
        self.assertEqual(len(analyze_syllabus('Course Syllabus\nWeek 1: Loads')['topics']), 1)
        indented = analyze_syllabus('Topics:\n  1. Loads\n\tTension and compression\n  2. Inspection')
        self.assertEqual([t['title'] for t in indented['topics']], ['Loads', 'Inspection'])
        self.assertEqual(indented['topics'][0]['subtopics'], ['Tension and compression'])
        with self.assertRaises(CourseError):
            build_course_markdown('Course Name: Structures\nDescription: A short course')

    def test_local_template_covers_topics_without_inventing_lesson_facts(self):
        content = build_course_markdown(SAMPLE)
        self.assertEqual(content.count('## Module '), 5)
        self.assertIn('[Add instructor-verified explanations', content)
        self.assertIn('Answer key', content)
        self.assertIn('Week(s):** 3-4', content)
        self.assertNotIn('## Module 1: Course Code', content)

    def test_markdown_bom_and_bad_uploads(self):
        self.assertEqual(extract_text('test.md', b'\xef\xbb\xbfWeek 1: Loads'), 'Week 1: Loads')
        for name, data in [('test.exe', b'file'), ('test.txt', b'\xff'), ('test.txt', b' '),
                           ('test.txt', b'a' * (MAX_FILE_BYTES + 1)), ('test.pdf', b'broken')]:
            with self.subTest(name=name, size=len(data)), self.assertRaises(CourseError):
                extract_text(name, data)

    def test_docx_body_order_preserved(self):
        from docx import Document
        document = Document()
        document.add_paragraph('Course Name: Structures')
        document.add_paragraph('Course Outline:')
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = 'Week', 'Topics'
        table.cell(1, 0).text, table.cell(1, 1).text = '1', 'Loads'
        document.add_paragraph('References:')
        document.add_paragraph('Approved handbook')
        target = io.BytesIO()
        document.save(target)
        plan = analyze_syllabus(extract_text('course.docx', target.getvalue()))
        self.assertEqual(plan['topics'][0]['title'], 'Loads')
        self.assertEqual(plan['references'], ['Approved handbook'])

    def test_text_pdf_and_scan_rejection(self):
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
        writer = PdfWriter()
        page = writer.add_blank_page(300, 300)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        stream = DecodedStreamObject()
        stream.set_data(b'BT /F1 12 Tf 20 200 Td (Week 1: Loads) Tj ET')
        page[NameObject('/Contents')] = stream
        target = io.BytesIO()
        writer.write(target)
        self.assertIn('Week 1: Loads', extract_text('course.pdf', target.getvalue()))
        writer = PdfWriter()
        writer.add_blank_page(100, 100)
        target = io.BytesIO()
        writer.write(target)
        with self.assertRaisesRegex(CourseError, 'OCR'):
            extract_text('scan.pdf', target.getvalue())

    def test_project_roundtrip_and_export_exclude_credentials(self):
        project = make_project(SAMPLE, 'Structures', build_course_markdown(SAMPLE))
        project['api_key'] = 'fake-key-never-save'
        self.assertNotIn(b'fake-key-never-save', project_bytes(project))
        restored = load_project(project_bytes(project))
        self.assertEqual(restored['content'], project['content'])
        with zipfile.ZipFile(io.BytesIO(export_bundle(project))) as archive:
            self.assertEqual(set(archive.namelist()), {'course.md', 'course.html', 'project.json', 'syllabus.txt', 'README.txt'})
            self.assertEqual(load_project(archive.read('project.json')), restored)

    def test_invalid_projects_do_not_load(self):
        for data in [b'bad', b'[]', b'{"schema_version":2}', b'{"schema_version":true}',
                     json.dumps({'schema_version': 1, 'title': 123}).encode()]:
            with self.subTest(data=data), self.assertRaises(CourseError):
                load_project(data)

    def test_atomic_save_preserves_old_file_on_failure(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'course.json'
            atomic_write(path, b'old')
            with patch('course_core.os.replace', side_effect=OSError('disk failure')):
                with self.assertRaises(OSError):
                    atomic_write(path, b'new')
            self.assertEqual(path.read_bytes(), b'old')
            self.assertEqual(list(Path(directory).iterdir()), [path])

    def test_html_export_escapes_active_content(self):
        data = export_html('<script>bad()</script>', '<img src=x onerror=bad()>').decode()
        self.assertNotIn('<script>', data)
        self.assertNotIn('<img ', data)
        self.assertIn('&lt;script&gt;', data)


if __name__ == '__main__':
    unittest.main()
